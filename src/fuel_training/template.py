"""The data-entry template for the field: what to record from now on, and why.

Every column exists because the exploration (docs/exploration.md) found the
model blind without it, or because the current sheets record it in a way that
cannot be trusted (free text for a number, planned and actual in one cell,
decimals exported as dates). One row per unit per day - the daily operation of
ADR 0001 - with the plan, what actually happened, and the fuel stick, side by side.

`fpt template` writes it; the unit and location lists come from production's
catalog so a name typed here resolves there.
"""

from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from fuel_training.catalog import Catalogs

DAY_STATUS = ["Operasi", "Standby", "Rusak", "Perbaikan", "Libur"]
ACTIVITY = [
    "Mobilisasi",
    "Lifting",
    "Mobilisasi + Lifting",
    "Angkut air (VT)",
    "Hisap limbah (VT)",
    "Lainnya",
]
ROAD = ["Aspal", "Tanah/berbatu", "Berlumpur", "Campuran"]
WEATHER = ["Cerah", "Hujan", "Hujan lebat"]


@dataclass(frozen=True, slots=True)
class Column:
    name: str
    kind: str  # tanggal | jam | angka | pilihan | teks
    unit: str
    required: str  # Wajib | Wajib bila ... | Opsional
    when: str  # Sebelum berangkat | Sesudah selesai
    why: str
    choices: tuple[str, ...] = ()
    example: object = ""


COLUMNS: tuple[Column, ...] = (
    # --- identity ------------------------------------------------------------------
    Column(
        "Tanggal",
        "tanggal",
        "",
        "Wajib",
        "Sebelum",
        "Kunci baris: satu kendaraan, satu hari.",
        example="2026-10-01",
    ),
    Column(
        "Kendaraan",
        "pilihan",
        "",
        "Wajib",
        "Sebelum",
        "Pilih dari daftar; nama bebas tidak dapat dicocokkan ke katalog.",
        example="VT 01",
    ),
    Column(
        "Status Hari",
        "pilihan",
        "",
        "Wajib",
        "Sesudah",
        "Standby/Rusak dipisah dari hari kerja; hari standby tetap membakar BBM (idle).",
        tuple(DAY_STATUS),
        "Operasi",
    ),
    Column(
        "ID Operasi (aplikasi)",
        "teks",
        "",
        "Opsional",
        "Sebelum",
        "Menghubungkan baris ini dengan estimasi yang dibuat di aplikasi.",
        example="OPR-...",
    ),
    # --- the plan ------------------------------------------------------------------
    Column(
        "Jenis Kegiatan",
        "pilihan",
        "",
        "Wajib bila Operasi",
        "Sebelum",
        "Kegiatan dari daftar, bukan dari teks bebas; hanya 3 crane yang boleh Lifting.",
        tuple(ACTIVITY),
        "Hisap limbah (VT)",
    ),
    Column(
        "Uraian Kegiatan",
        "teks",
        "",
        "Opsional",
        "Sebelum",
        "Konteks untuk manusia; model tidak membacanya.",
        example="Hisap limbah di SP-II, buang ke WS Limau",
    ),
    Column(
        "Peminta (User)",
        "teks",
        "",
        "Opsional",
        "Sebelum",
        "Pola per peminta/lokasi kerja.",
        example="Pak Rais",
    ),
    Column(
        "Rute (urutan lokasi)",
        "teks",
        "",
        "Wajib bila Operasi",
        "Sebelum",
        "Nama dari Daftar Lokasi dipisah ' - ', termasuk kembali ke pool. Dipakai untuk jarak rute & riwayat serupa.",
        example="POOL LIMAU - SP-II - WS LIMAU - POOL LIMAU",
    ),
    Column(
        "Jarak Rencana (km)",
        "angka",
        "km",
        "Wajib bila Operasi",
        "Sebelum",
        "Jarak yang direncanakan; inilah yang diketahui saat memprediksi.",
        example=62.5,
    ),
    Column(
        "Jam Lifting Rencana",
        "angka",
        "jam",
        "Wajib bila Lifting",
        "Sebelum",
        "Hanya Truck Crane 01/02 dan Wheel Crane.",
        example="",
    ),
    Column(
        "Jumlah Ritase",
        "angka",
        "trip",
        "Wajib untuk VT",
        "Sebelum",
        "Jumlah angkut/hisap bolak-balik; tiap trip ada muat/bongkar.",
        example=3,
    ),
    Column(
        "Jenis Muatan",
        "teks",
        "",
        "Opsional",
        "Sebelum",
        "Pipa, alat berat, air bersih, limbah…",
        example="Limbah cair",
    ),
    Column(
        "Berat / Volume Muatan",
        "angka",
        "ton atau kL",
        "Wajib bila ada muatan",
        "Sebelum",
        "Beban paling menentukan konsumsi; sekarang tidak tercatat sama sekali.",
        example=12,
    ),
    Column(
        "BBM Disiapkan (L)",
        "angka",
        "liter",
        "Wajib bila BBM diberikan",
        "Sebelum",
        "Liter yang diberikan hari itu untuk kendaraan ini (bukan isi ulang beberapa hari).",
        example=60,
    ),
    # --- what actually happened -------------------------------------------------------
    Column(
        "Jam Mulai",
        "jam",
        "",
        "Wajib bila Operasi",
        "Sesudah",
        "Lama operasi dan jam kerja.",
        example="07:30",
    ),
    Column("Jam Selesai", "jam", "", "Wajib bila Operasi", "Sesudah", "", example="16:00"),
    Column(
        "Odometer Awal (km)",
        "angka",
        "km",
        "Wajib",
        "Sebelum",
        "Jarak aktual dari odometer; pembanding GPS.",
        example=12345.6,
    ),
    Column("Odometer Akhir (km)", "angka", "km", "Wajib", "Sesudah", "", example=12404.1),
    Column(
        "Hour Meter Awal (jam)",
        "angka",
        "jam",
        "Wajib",
        "Sebelum",
        "Jam mesin menyala, termasuk idle dan pompa; menjelaskan BBM saat tidak bergerak.",
        example=4521.2,
    ),
    Column("Hour Meter Akhir (jam)", "angka", "jam", "Wajib", "Sesudah", "", example=4529.0),
    Column(
        "Jam PTO / Pompa (jam)",
        "angka",
        "jam",
        "Wajib untuk VT",
        "Sesudah",
        "Pompa vakum memakai BBM tanpa bergerak.",
        example=2.5,
    ),
    Column(
        "Jam Lifting Aktual",
        "angka",
        "jam",
        "Wajib bila Lifting",
        "Sesudah",
        "Dari hour meter crane, bukan rencana: rencana vs aktual tidak boleh satu kolom.",
        example="",
    ),
    Column(
        "Kondisi Jalan",
        "pilihan",
        "",
        "Opsional",
        "Sesudah",
        "Lumpur menaikkan konsumsi.",
        tuple(ROAD),
        "Tanah/berbatu",
    ),
    Column(
        "Cuaca",
        "pilihan",
        "",
        "Opsional",
        "Sesudah",
        "Hujan memperlambat dan menaikkan konsumsi.",
        tuple(WEATHER),
        "Cerah",
    ),
    Column(
        "Pengemudi / Operator",
        "teks",
        "",
        "Opsional",
        "Sesudah",
        "Gaya mengemudi berpengaruh; gunakan nama/ID yang sama setiap kali.",
        example="OP-017",
    ),
    # --- the fuel stick ---------------------------------------------------------------
    Column(
        "Stick Awal (L)",
        "angka",
        "liter",
        "Wajib",
        "Sebelum",
        "Isi tangki sebelum berangkat.",
        example=180.5,
    ),
    Column(
        "Isi Ulang (L)",
        "angka",
        "liter",
        "Wajib (0 bila tidak ada)",
        "Sesudah",
        "Setiap isi ulang hari itu, dijumlah.",
        example=0,
    ),
    Column(
        "Drain (L)",
        "angka",
        "liter",
        "Wajib (0 bila tidak ada)",
        "Sesudah",
        "BBM yang dikeluarkan dari tangki bukan untuk mesin.",
        example=0,
    ),
    Column(
        "Stick Akhir (L)",
        "angka",
        "liter",
        "Wajib",
        "Sesudah",
        "Isi tangki setelah selesai. Konsumsi dihitung dari keempat angka ini; jangan diketik manual.",
        example=141.3,
    ),
    Column(
        "Catatan",
        "teks",
        "",
        "Opsional",
        "Sesudah",
        "Kejadian yang menjelaskan angka aneh: macet, rusak di jalan, idle lama, stick tidak terbaca.",
        example="",
    ),
)

INSTRUCTIONS = [
    "Satu baris = satu kendaraan, satu hari. Hari Standby/Rusak/Libur tetap diisi (Status Hari).",
    "Kolom angka hanya berisi angka. Jangan menulis '80.8.', '2 jam', atau teks. Desimal memakai titik.",
    "Jangan mengetik hasil perhitungan (konsumsi, km/L): isi angka mentahnya, sistem yang menghitung.",
    "Rencana dan aktual ditulis di kolom yang berbeda. Jangan menimpa rencana dengan aktual.",
    "Kendaraan, Jenis Kegiatan, Status Hari, Kondisi Jalan, Cuaca dipilih dari daftar.",
    "Rute memakai nama dari sheet 'Daftar Lokasi', dipisah ' - ', berurutan, termasuk kembali ke pool.",
    "BBM Disiapkan = liter yang diberikan untuk hari ini. Bila tangki diisi untuk beberapa hari, catat di Isi Ulang, bukan di BBM Disiapkan.",
    "Bila Google Sheets: format kolom angka sebagai 'Angka' (Format > Angka) agar tidak berubah menjadi tanggal saat diekspor.",
]


def write_template(path: Path, catalogs: Catalogs) -> Path:
    workbook = Workbook()
    header_fill = PatternFill("solid", fgColor="1C5871")
    header_font = Font(bold=True, color="FFFFFF")

    guide = workbook.active
    guide.title = "Petunjuk"
    guide["A1"] = "Petunjuk pengisian data operasi harian"
    guide["A1"].font = Font(bold=True, size=14)
    for index, line in enumerate(INSTRUCTIONS, start=3):
        guide.cell(row=index, column=1, value=f"{index - 2}. {line}")
    guide.column_dimensions["A"].width = 130

    vehicles = workbook.create_sheet("Daftar Kendaraan")
    for position, title in enumerate(("Kendaraan", "Tipe", "Grup", "Alias"), start=1):
        cell = vehicles.cell(row=1, column=position, value=title)
        cell.fill, cell.font = header_fill, header_font
    for row, option in enumerate(catalogs.vehicles.options(), start=2):
        vehicles.cell(row=row, column=1, value=option.name)
        vehicles.cell(row=row, column=2, value=option.type)
        vehicles.cell(row=row, column=3, value=option.group)
        vehicles.cell(row=row, column=4, value=", ".join(option.aliases))
    last_vehicle = len(catalogs.vehicles.options()) + 1
    for letter, width in zip("ABCD", (24, 18, 16, 30), strict=True):
        vehicles.column_dimensions[letter].width = width

    places = workbook.create_sheet("Daftar Lokasi")
    places.cell(row=1, column=1, value="Lokasi").fill = header_fill
    places["A1"].font = header_font
    for row, place in enumerate(catalogs.locations.options(), start=2):
        places.cell(row=row, column=1, value=place.name)
    places.column_dimensions["A"].width = 40

    lists = workbook.create_sheet("Pilihan")
    for position, values in enumerate((DAY_STATUS, ACTIVITY, ROAD, WEATHER), start=1):
        for row, value in enumerate(values, start=1):
            lists.cell(row=row, column=position, value=value)
    lists.sheet_state = "hidden"
    list_ranges = {
        tuple(DAY_STATUS): f"Pilihan!$A$1:$A${len(DAY_STATUS)}",
        tuple(ACTIVITY): f"Pilihan!$B$1:$B${len(ACTIVITY)}",
        tuple(ROAD): f"Pilihan!$C$1:$C${len(ROAD)}",
        tuple(WEATHER): f"Pilihan!$D$1:$D${len(WEATHER)}",
    }

    entry = workbook.create_sheet("Operasi Harian", 1)
    rows = 2000
    for index, column in enumerate(COLUMNS, start=1):
        letter = get_column_letter(index)
        cell = entry.cell(row=1, column=index, value=column.name)
        cell.fill, cell.font = header_fill, header_font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        entry.column_dimensions[letter].width = max(14, min(34, len(column.name) + 4))
        example = entry.cell(
            row=2, column=index, value=column.example if column.example != "" else None
        )
        example.font = Font(italic=True, color="888888")
        target = f"{letter}2:{letter}{rows}"
        if column.name == "Kendaraan":
            rule = DataValidation(
                type="list", formula1=f"'Daftar Kendaraan'!$A$2:$A${last_vehicle}"
            )
        elif column.choices:
            rule = DataValidation(type="list", formula1=list_ranges[column.choices])
        elif column.kind == "angka":
            rule = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0")
            rule.error = "Isi dengan angka (desimal memakai titik), tanpa satuan atau teks."
            for row in range(2, rows + 1):
                entry.cell(row=row, column=index).number_format = "0.0##"
        elif column.kind == "tanggal":
            rule = DataValidation(type="date", operator="greaterThan", formula1="44927")
            rule.error = "Isi dengan tanggal."
            for row in range(2, rows + 1):
                entry.cell(row=row, column=index).number_format = "yyyy-mm-dd"
        elif column.kind == "jam":
            rule = DataValidation(type="time")
            rule.error = "Isi dengan jam, misalnya 07:30."
            for row in range(2, rows + 1):
                entry.cell(row=row, column=index).number_format = "hh:mm"
        else:
            continue
        rule.showErrorMessage = True
        rule.add(target)
        entry.add_data_validation(rule)
    entry.freeze_panes = "C2"
    entry.row_dimensions[1].height = 45

    dictionary = workbook.create_sheet("Kamus Kolom", 2)
    headers = ("Kolom", "Jenis", "Satuan", "Wajib?", "Diisi", "Kenapa dibutuhkan")
    for position, title in enumerate(headers, start=1):
        cell = dictionary.cell(row=1, column=position, value=title)
        cell.fill, cell.font = header_fill, header_font
    for row, column in enumerate(COLUMNS, start=2):
        cells = (column.name, column.kind, column.unit, column.required, column.when, column.why)
        for index, value in enumerate(cells, start=1):
            cell = dictionary.cell(row=row, column=index, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for letter, width in zip("ABCDEF", (26, 10, 12, 24, 12, 90), strict=True):
        dictionary.column_dimensions[letter].width = width

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
